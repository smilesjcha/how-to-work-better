"""PPTX 무결성·시간·렌더된 글자의 텍스트 상자 내부 범위를 검사한다.

자동 검사는 이미지 확대 검수를 대체하지 않는다.
"""
import json
import unicodedata
from pathlib import Path

import pdfplumber
from pptx import Presentation
from slides_data import SLIDES, BLOCKS

ROOT = Path(__file__).resolve().parents[1]
PT_PER_EMU = 72 / 914400
TOLERANCE_PT = 2.0


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFKC', value) if not c.isspace())


def text_targets(slide):
    for si, shape in enumerate(slide.shapes):
        if shape.has_text_frame and shape.text.strip():
            yield f'shape-{si}', shape.text, tuple(v * PT_PER_EMU for v in (shape.left, shape.top, shape.left+shape.width, shape.top+shape.height))
        if shape.has_table:
            table = shape.table
            y = shape.top
            for ri, row in enumerate(table.rows):
                x = shape.left
                for ci, col in enumerate(table.columns):
                    cell = table.cell(ri, ci)
                    if cell.text.strip():
                        yield f'table-{si}-{ri}-{ci}', cell.text, tuple(v * PT_PER_EMU for v in (x,y,x+col.width,y+row.height))
                    x += col.width
                y += row.height


def main():
    pptx = ROOT/'output/samcheok-ai-docs-workshop.pptx'
    pdf = ROOT/'output/samcheok-ai-docs-workshop.pdf'
    prs = Presentation(pptx)
    assert len(prs.slides) == len(SLIDES)
    assert len(BLOCKS) == 18
    assert all(b['minutes'] == 10 for b in BLOCKS)
    assert sum(d['minutes'] for d in SLIDES) == 180
    for block in BLOCKS:
        assert sum(d['minutes'] for d in SLIDES if d['block'] == block['key']) == block['minutes']
    results = {'slides':len(SLIDES),'minutes':180,'matched_text_objects':0,'unmatched':[],'overflow':[],'canvas_outside':[]}
    with pdfplumber.open(pdf) as doc:
        assert len(doc.pages) == len(prs.slides)
        for n,(slide,page) in enumerate(zip(prs.slides,doc.pages),1):
            characters=[]
            for c in page.chars:
                for letter in normalized(c['text']): characters.append((letter,c))
            content=''.join(c[0] for c in characters)
            for name,value,box in text_targets(slide):
                wanted=normalized(value)
                matches=[]; pos=content.find(wanted)
                while pos>=0:
                    coords=[c[1] for c in characters[pos:pos+len(wanted)]]
                    bounds=(min(c['x0'] for c in coords),min(c['top'] for c in coords),max(c['x1'] for c in coords),max(c['bottom'] for c in coords))
                    distance=abs(bounds[0]-box[0])+abs(bounds[1]-box[1])
                    matches.append((distance,bounds)); pos=content.find(wanted,pos+1)
                if not matches:
                    results['unmatched'].append({'slide':n,'name':name,'text':value[:70]}); continue
                results['matched_text_objects']+=1
                bounds=min(matches)[1]
                if bounds[0]<box[0]-TOLERANCE_PT or bounds[1]<box[1]-TOLERANCE_PT or bounds[2]>box[2]+TOLERANCE_PT or bounds[3]>box[3]+TOLERANCE_PT:
                    results['overflow'].append({'slide':n,'name':name,'text':value[:70],'box':box,'rendered':bounds})
            for shape in slide.shapes:
                if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width+10 or shape.top+shape.height>prs.slide_height+10:
                    results['canvas_outside'].append({'slide':n,'name':shape.name})
    out=ROOT/'tmp/qa-text-bounds.json'; out.parent.mkdir(exist_ok=True); out.write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False,indent=2))
    if results['overflow'] or results['canvas_outside'] or results['unmatched']: raise SystemExit(1)


if __name__ == '__main__': main()
