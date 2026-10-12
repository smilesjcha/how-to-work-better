"""편집형 장표 레이아웃. 디자인 값은 theme.py에서만 관리한다."""
import os

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.image import Image

import theme as T

TOTAL = 0
ASSETS = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "assets"))


def font(run, size=T.PT_BODY, color=T.INK, bold=False, mono=False):
    run.font.name = T.FONT_MONO if mono else T.FONT
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color


def paragraphs(tf, value, size, color, bold, align, spacing, mono=False):
    """빈 문단의 간격도 명시해서 테마의 기본 문단 간격을 차단한다."""
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = T.ZERO
        p.space_after = T.ZERO
        p.font.name = T.FONT
        p.font.size = size if line else T.PT_SPACER
        run = p.add_run()
        run.text = line
        font(run, size if line else T.PT_SPACER, color, bold, mono)


def text(slide, x, y, w, h, value, size=T.PT_BODY, color=T.INK,
         bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, mono=False):
    box = slide.shapes.add_textbox(*(int(v) for v in (x, y, w, h)))
    tf = box.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = T.ZERO
    paragraphs(tf, value, size, color, bold, align, T.TEXT_LINE_SPACING, mono)
    return box


def clear_effects(shape):
    """Office 테마가 직사각형·화살표에 기본 그림자를 적용하지 않게 한다."""
    for effect in shape._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    props = shape._element.spPr
    for child in list(props):
        if child.tag.endswith("}effectLst") or child.tag.endswith("}effectDag"):
            props.remove(child)
    props.append(OxmlElement("a:effectLst"))


def rect(slide, x, y, w, h, fill=T.WHITE, line=None, rounded=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, *(int(v) for v in (x, y, w, h)))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = T.LINE_THIN
    if rounded:
        shape.adjustments[0] = T.BOX_RADIUS
    clear_effects(shape)
    return shape


def shape_text(shape, value, size=T.PT_BODY, color=T.INK, bold=False,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
               pad_x=T.PAD, pad_y=T.PAD):
    tf = shape.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = pad_x
    tf.margin_top = tf.margin_bottom = pad_y
    paragraphs(tf, value, size, color, bold, align, T.BOX_LINE_SPACING)


def rule(slide, x, y, width, color=T.LINE, weight=T.LINE_THIN):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(int(v) for v in (x, y, x + width, y)))
    line.line.color.rgb = color
    line.line.width = weight
    clear_effects(line)


def base(prs, data, bg=T.PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide


def frame(slide, data, dark=False):
    c = T.WHITE if dark else T.INK
    fc = T.TEAL_LIGHT if dark else T.MUTED
    if not dark:
        text(slide, *T.TITLE_AREA, data["title"], T.PT_TITLE, c, True)
        rule(slide, T.MX, T.RULE_Y, T.RULE_W, T.TEAL, T.LINE_MED)
    footer = f"{data['time_start']}-{data['time_end']} / 삼척새일센터 / 생성형 AI 실무 문서"
    text(slide, *T.FOOTER_AREA, footer, T.PT_CAPTION, fc)
    text(slide, *T.PAGE_AREA, f"{data['no']:02d} / {TOTAL:02d}", T.PT_CAPTION, fc, align=PP_ALIGN.RIGHT)
    notes = slide.notes_slide.notes_text_frame
    notes.text = data.get("note", "")
    for p in notes.paragraphs:
        for r in p.runs:
            font(r, T.PT_TINY, T.INK)
    return slide


def source_caption(slide, d):
    if d.get("source_file"):
        box = text(slide, *T.SOURCE_FILE_AREA, d["source_file"], T.PT_CAPTION, T.MUTED)
        if d.get("download_url"):
            for p in box.text_frame.paragraphs:
                for r in p.runs:
                    r.hyperlink.address = d["download_url"]


def cover(prs, d):
    s = base(prs, d, T.NAVY)
    text(s, *T.COVER_KICKER_AREA, "삼척새일센터 실무 문서 작성 워크숍", T.PT_SMALL, T.TEAL_LIGHT, True)
    text(s, *T.COVER_TITLE_AREA, d["title"], T.PT_COVER, T.WHITE, True)
    text(s, *T.COVER_SUBTITLE_AREA, d["subtitle"], T.PT_LEAD, T.APRICOT)
    text(s, *T.COVER_DATE_AREA, "2026. 10. 13. (화) 14:00–17:00 · 총 180분", T.PT_BODY, T.WHITE)
    text(s, *T.COVER_AUTHOR_AREA, "강사 차성재 · 무신사 Core AI Product Manager · 아주대학교 AI대학원 겸임교수", T.PT_SMALL, T.TEAL_LIGHT)
    return frame(s, d, True)


def section(prs, d):
    s = base(prs, d, T.NAVY)
    text(s, *T.SECTION_KICKER_AREA, d.get("kicker", ""), T.PT_SMALL, T.APRICOT, True, PP_ALIGN.CENTER)
    text(s, *T.SECTION_TITLE_AREA, d["title"], T.PT_COVER, T.WHITE, True, PP_ALIGN.CENTER)
    text(s, *T.SECTION_SUBTITLE_AREA, d.get("sub", ""), T.PT_BODY, T.TEAL_LIGHT, align=PP_ALIGN.CENTER)
    return frame(s, d, True)


def statement(prs, d):
    s = frame(base(prs, d), d)
    text(s, *T.STATEMENT_LEAD_AREA, d["lead"], T.PT_LEAD, T.INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, *T.STATEMENT_CAPTION_AREA, d.get("caption", ""), T.PT_BODY, T.MUTED)
    return s


def bullets(prs, d):
    s = frame(base(prs, d), d)
    text(s, *T.BULLETS_LEAD_AREA, d.get("lead", ""), T.PT_LEAD, T.TEAL, True)
    text(s, *T.BULLETS_BODY_AREA, "\n\n".join(f"{i + 1}.  {v}" for i, v in enumerate(d["items"])), T.PT_BODY)
    source_caption(s, d)
    return s


def schedule(prs, d):
    return grid(prs, d, ["시간", "분", "내용"], d["rows"], T.SCHEDULE_WIDTHS)


def flow(prs, d):
    s = frame(base(prs, d), d)
    labels = d["items"]
    width = (T.CW - T.GAP * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        x = T.MX + i * (width + T.GAP)
        focused = i == d.get("focus", -1)
        shape = rect(s, x, T.FLOW_Y, width, T.FLOW_H, T.TEAL_LIGHT if focused else T.WHITE, T.TEAL, True)
        shape_text(shape, label, T.PT_BODY, T.INK, focused, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < len(labels) - 1:
            arrow = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, *(int(v) for v in (x + width, T.FLOW_ARROW_Y, T.GAP, T.FLOW_ARROW_H)))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = T.TEAL
            arrow.line.fill.background()
            clear_effects(arrow)
    text(s, *T.FLOW_CAPTION_AREA, d.get("caption", ""), T.PT_BODY, T.MUTED)
    return s


def columns(prs, d):
    s = frame(base(prs, d), d)
    for i in range(len(d["titles"])):
        x = T.MX + i * (T.HALF_W + T.GAP)
        text(s, x, T.COLUMN_TITLE_Y, T.HALF_W, T.COLUMN_TITLE_H, d["titles"][i], T.PT_BODY, T.TEAL if i == 0 else T.INK, True)
        rule(s, x, T.COLUMN_RULE_Y, T.HALF_W)
        text(s, x, T.COLUMN_BODY_Y, T.HALF_W, T.COLUMN_BODY_H, d["bodies"][i], T.PT_BODY)
    return s


def prompt(prs, d):
    s = frame(base(prs, d), d)
    text(s, *T.PROMPT_INSTRUCTION_AREA, d.get("instruction", "그대로 복사하여 붙여넣으세요"), T.PT_SMALL, T.TEAL, True)
    shape = rect(s, *T.PROMPT_AREA, T.MIST, T.LINE, True)
    shape_text(shape, d["text"], T.PT_CODE, T.INK)
    source_caption(s, d)
    return s


def document(prs, d):
    s = frame(base(prs, d), d)
    shape = rect(s, *T.DOCUMENT_AREA, T.WHITE, T.LINE)
    shape_text(shape, d["doc"], T.PT_DOCUMENT)
    if d.get("label"):
        is_error = d["label"] == "의도적 오류"
        fill = T.ERROR if is_error else (T.APRICOT if d["label"] == "가정" else T.TEAL_LIGHT)
        label = rect(s, *T.DOCUMENT_LABEL_AREA, fill)
        shape_text(label, d["label"], T.PT_CAPTION, T.WHITE if is_error else T.INK, True,
                   PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, T.LABEL_PAD_X, T.LABEL_PAD_Y)
    source_caption(s, d)
    return s


def grid(prs, d, headers, rows, widths=None):
    if any(len(row) != len(headers) for row in rows):
        raise ValueError(f"표의 열 수가 일치하지 않습니다: {d.get('title', '')}")
    s = frame(base(prs, d), d)
    table = s.shapes.add_table(len(rows) + 1, len(headers), *(int(v) for v in T.TABLE_AREA)).table
    for i, width in enumerate(widths or [T.CW / len(headers)] * len(headers)):
        table.columns[i].width = int(width)
    for ri, row in enumerate([headers] + rows):
        for ci, value in enumerate(row):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = T.NAVY if ri == 0 else (T.MIST if ri % 2 else T.WHITE)
            if d.get('type') == 'gantt' and ri > 0 and ci > 0:
                cell.fill.fore_color.rgb = T.TEAL_LIGHT if value else T.WHITE
            tf = cell.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = T.TABLE_PAD_X
            tf.margin_top = tf.margin_bottom = T.TABLE_PAD_Y
            paragraphs(tf, value, T.PT_TABLE, T.WHITE if ri == 0 else T.INK,
                       ri == 0, PP_ALIGN.LEFT, T.TABLE_LINE_SPACING)
    if d.get("caption"):
        text(s, *T.TABLE_CAPTION_AREA, d["caption"], T.PT_CAPTION, T.MUTED)
    source_caption(s, d)
    return s


def table(prs, d):
    return grid(prs, d, d["headers"], d["rows"])


def gantt(prs, d):
    return grid(prs, d, d["headers"], d["rows"], T.GANTT_WIDTHS)


def quiz(prs, d):
    s = frame(base(prs, d, T.MIST), d)
    text(s, *T.QUIZ_QUESTION_AREA, d["question"], T.PT_LEAD, T.INK, True)
    for i, option in enumerate(d["options"]):
        x = T.MX + i * (T.HALF_W + T.GAP)
        shape = rect(s, x, T.QUIZ_OPTION_Y, T.HALF_W, T.QUIZ_OPTION_H, T.WHITE, T.LINE, True)
        shape_text(shape, option, T.PT_BODY, T.INK, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, *T.QUIZ_ANSWER_AREA, d["answer"], T.PT_BODY, T.TEAL, True)
    return s


def fitted_picture(slide, path, area):
    """이미지 전체를 토큰 영역 안에 종횡비를 유지해 넣는다."""
    x, y, box_w, box_h = area
    pixel_w, pixel_h = Image.from_file(path).size
    scale = min(box_w / pixel_w, box_h / pixel_h)
    width, height = pixel_w * scale, pixel_h * scale
    return slide.shapes.add_picture(path, int(x + (box_w - width) / 2), int(y + (box_h - height) / 2), width=int(width), height=int(height))


def image(prs, d):
    s = frame(base(prs, d), d)
    path = os.path.join(ASSETS, d["path"])
    if os.path.exists(path):
        fitted_picture(s, path, T.IMAGE_AREA)
    text(s, *T.IMAGE_CAPTION_AREA, d["caption"], T.PT_BODY, T.MUTED, align=PP_ALIGN.CENTER)
    return s


def photo(prs, d):
    s = frame(base(prs, d), d)
    path = os.path.join(ASSETS, d["path"])
    if os.path.exists(path):
        fitted_picture(s, path, T.PHOTO_AREA)
    text(s, *T.PHOTO_LEAD_AREA, d["lead"], T.PT_LEAD, T.INK, True)
    text(s, *T.PHOTO_CAPTION_AREA, d["caption"], T.PT_BODY, T.MUTED)
    source_caption(s, d)
    return s


def document_preview(prs, d):
    s = frame(base(prs, d), d)
    path = os.path.join(ASSETS, d["path"])
    if os.path.exists(path):
        fitted_picture(s, path, T.PREVIEW_IMAGE_AREA)
    text(s, *T.PREVIEW_LEAD_AREA, d["lead"], T.PT_BODY, T.TEAL, True)
    text(s, *T.PREVIEW_CAPTION_AREA, d["caption"], T.PT_BODY, T.INK)
    source_caption(s, d)
    return s


def pause(prs, d):
    s = base(prs, d, T.TEAL)
    text(s, *T.PAUSE_TITLE_AREA, d["title"], T.PT_COVER, T.WHITE, True, PP_ALIGN.CENTER)
    text(s, *T.PAUSE_MESSAGE_AREA, d["message"], T.PT_BODY, T.TEAL_LIGHT, align=PP_ALIGN.CENTER)
    return frame(s, d, True)


def profile(prs, d):
    s = frame(base(prs, d), d)
    path = os.path.join(ASSETS, "profile_photo.jpg")
    if os.path.exists(path):
        pixel_w, pixel_h = Image.from_file(path).size
        cropped_ratio = pixel_w * (1 - T.PROFILE_CROP_LEFT) / pixel_h
        portrait = s.shapes.add_picture(path, int(T.PROFILE_PHOTO_X), int(T.PROFILE_PHOTO_Y),
                                       width=int(T.PROFILE_PHOTO_H * cropped_ratio), height=int(T.PROFILE_PHOTO_H))
        portrait.crop_left = T.PROFILE_CROP_LEFT
    text(s, *T.PROFILE_NAME_AREA, "차성재", T.PT_LEAD, T.INK, True)
    text(s, *T.PROFILE_ROLE_AREA, "무신사 Core AI AI Product Manager\n아주대학교 AI대학원 겸임교수", T.PT_BODY, T.TEAL, True)
    text(s, *T.PROFILE_BIO_AREA, "AI 제품을 기획하고 만들고 운영해온 실무형 AI PM\n금융의 표, 의료의 영상, 교육의 언어 데이터를 다뤘습니다. 지금은 이커머스에서 고객과 상품을 연결하는 Agent 제품을 만들고 있습니다.", T.PT_BODY)
    return s


def closing(prs, d):
    s = base(prs, d, T.NAVY)
    text(s, *T.CLOSING_TITLE_AREA, d["title"], T.PT_COVER, T.WHITE, True, PP_ALIGN.CENTER)
    text(s, *T.CLOSING_MESSAGE_AREA, d["message"], T.PT_BODY, T.APRICOT, align=PP_ALIGN.CENTER)
    return frame(s, d, True)


TYPE_TO_FN = {name: globals()[name] for name in "cover section statement bullets schedule flow columns prompt document table gantt quiz image photo document_preview pause profile closing".split()}
