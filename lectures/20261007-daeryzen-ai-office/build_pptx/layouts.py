"""편집형 강의 장표 렌더러. 색·서체·간격은 theme.py에서만 관리한다."""
import os
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.parts.image import Image
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont
from functools import lru_cache
import theme as T

TOTAL = 0
ASSETS = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'assets'))


@lru_cache(maxsize=16)
def _measure_font(size):
    return ImageFont.truetype(T.FONT_PATH, round(size.pt*T.MEASURE_SCALE))


def wrap_text(value, width, size):
    """실측 폭으로 줄바꿈. 렌더러 차이를 흡수할 안전 여유를 둔다."""
    font = _measure_font(size)
    limit = width/12700*T.MEASURE_SCALE*T.WRAP_SAFE
    lines = []
    for paragraph in str(value).split('\n'):
        current = ''
        for char in paragraph:
            if current and font.getlength(current+char) > limit:
                # 가능한 한 어절 단위로 나누되 긴 한 단어는 폭을 우선한다.
                if ' ' in current and font.getlength(current.rsplit(' ', 1)[0]) > limit/2:
                    before, after = current.rsplit(' ', 1)
                    lines.append(before.rstrip()); current = after+char
                else:
                    lines.append(current.rstrip()); current = char.lstrip()
            else:
                current += char
        lines.append(current.rstrip())
    return '\n'.join(lines)


def prompt_rows(items, width):
    """여러 줄 요청문에 필요한 높이를 배분해 구분선과 충돌을 막는다."""
    wrapped = [wrap_text(item, width, T.PROMPT_TEXT) for item in items]
    heights = [max(T.PROMPT_MIN_H,
                   (text.count('\n')+1)*T.PROMPT_TEXT*T.TEXT_LEADING+2*T.PROMPT_PAD_Y)
               for text in wrapped]
    if sum(heights) > T.STAGE_H:
        raise ValueError('프롬프트가 본문 높이를 초과합니다. 요청 항목을 분할하세요.')
    return list(zip(wrapped, heights))


def textbox(slide, x, y, w, h, value, size=T.BODY, color=T.INK, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, mono=False):
    shape = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for index, part in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = T.BODY_LINES
        p.space_before = p.space_after = 0
        run = p.add_run(); run.text = part
        run.font.name = T.MONO if mono else (T.FONT_HEAVY if bold else T.FONT)
        run.font.size = size; run.font.bold = bold; run.font.color.rgb = color
    return shape


def _clean(shape):
    for effect in shape._element.xpath('.//a:effectRef'):
        effect.set('idx', '0')
    sp = shape._element.spPr
    for child in list(sp):
        if child.tag.endswith('}effectLst') or child.tag.endswith('}effectDag'):
            sp.remove(child)
    sp.append(OxmlElement('a:effectLst'))


def rect(slide, x, y, w, h, fill=T.SURFACE, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, int(x), int(y), int(w), int(h))
    if rounded: shape.adjustments[0] = T.CORNER
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.fill.background(); _clean(shape)
    return shape


def line(slide, x, y, w, color=T.LINE, thick=T.STROKE):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       int(x), int(y), int(x+w), int(y))
    shape.line.color.rgb = color; shape.line.width = thick; _clean(shape)


def vertical(slide, x, y, h, color=T.LINE, thick=T.STROKE):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       int(x), int(y), int(x), int(y+h))
    shape.line.color.rgb = color; shape.line.width = thick; _clean(shape)


def base(prs, d, dark=False, title=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = T.INK if dark else T.PAPER
    if title:
        size = T.TITLE if len(d['title']) < 36 else T.SMALL_NUM
        textbox(s, T.MX, T.HEADER_Y, T.CW, T.HEADER_H, d['title'], size,
                T.WHITE if dark else T.INK, True)
        line(s, T.MX, T.RULE_Y, T.CW, T.GRAY_DARK if dark else T.LINE)
        rect(s, T.MX, T.RULE_Y-T.RULE_THICK/2, T.RULE_W, T.RULE_THICK,
             T.WARM if dark else T.GOOD)
    foot = T.PALE if dark else T.MUTED
    textbox(s, T.MX, T.FOOTER_Y, T.CW, T.FOOTER_H,
            d.get('phase', 'DAERYZEN · AI OFFICE'), T.CAPTION, foot)
    textbox(s, T.W-T.MX-T.RULE_W, T.FOOTER_Y, T.RULE_W, T.FOOTER_H,
            f"{d['no']:03d} / {TOTAL:03d}", T.PAGE, foot, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = d.get('note', '')
    return s


def lead(s, d):
    if d.get('lead'):
        textbox(s, T.MX, T.LEAD_Y, T.CW, T.LEAD_H, d['lead'], T.LEAD, T.MUTED)


def bottom(s, d):
    if d.get('bottom'):
        line(s, T.MX, T.STAGE_BOTTOM+T.RULE_SPACING, T.CW)
        textbox(s, T.MX, T.STAGE_BOTTOM+T.EDITORIAL_INSET/2,
                T.CW, T.FOOTER_H*2, d['bottom'], T.CAPTION, T.MUTED)


def cover(prs, d):
    s = base(prs, d, dark=True, title=False)
    vertical(s, T.MX, T.COVER_KICKER_Y, T.COVER_META_Y-T.COVER_KICKER_Y,
             T.WARM, T.RULE_THICK)
    x = T.MX+T.EDITORIAL_INSET
    textbox(s, x, T.COVER_KICKER_Y, T.CW, T.PANEL_TITLE_H,
            'DAERYZEN  /  OFFICE WORKSHOP', T.KICKER, T.WARM, True)
    textbox(s, x, T.COVER_TITLE_Y, T.CW-T.EDITORIAL_INSET,
            T.SECTION_TITLE_Y, d['title'], T.COVER, T.WHITE, True)
    textbox(s, x, T.COVER_SUB_Y, T.CW-T.EDITORIAL_INSET,
            T.LEAD_H*2, d['lead'], T.LEAD, T.PALE)
    line(s, x, T.COVER_META_Y-T.GAP, T.CW-T.EDITORIAL_INSET, T.GRAY_DARK)
    textbox(s, x, T.COVER_META_Y, T.CW-T.EDITORIAL_INSET,
            T.LEAD_H, d.get('bottom', ''), T.BODY_SM, T.WHITE)
    return s


def section(prs, d):
    s = base(prs, d, dark=True, title=False)
    textbox(s, T.MX, T.SECTION_NO_Y, T.CW, T.PANEL_TITLE_H,
            d.get('eyebrow', ''), T.KICKER, T.WARM, True)
    line(s, T.MX, T.SECTION_TITLE_Y-T.GAP, T.CW, T.GRAY_DARK)
    textbox(s, T.MX, T.SECTION_TITLE_Y, T.CW, T.LEAD_H*3,
            d['title'], T.COVER, T.WHITE, True)
    textbox(s, T.MX, T.SECTION_SUB_Y+T.LEAD_H, T.CW, T.LEAD_H*2,
            d.get('lead', ''), T.BODY, T.PALE)
    return s


def triad(prs, d):
    s = base(prs, d); lead(s, d)
    mode = d['_composition']; items = d['cards']
    if mode == 'index':
        height = T.STAGE_H/3
        for i, (head, body) in enumerate(items):
            y = T.STAGE_Y+i*height
            line(s, T.MX, y, T.CW)
            textbox(s, T.MX, y+T.PAD, T.ROW_NUM_W, height-T.PAD,
                    f'{i+1:02d}', T.SMALL_NUM, T.GOOD, True)
            textbox(s, T.MX+T.ROW_NUM_W, y+T.PAD,
                    T.WIDE_LEFT-T.ROW_NUM_W, height-T.PAD,
                    head, T.BODY, T.INK, True, anchor=MSO_ANCHOR.MIDDLE)
            textbox(s, T.MX+T.WIDE_LEFT+T.GAP, y+T.PAD,
                    T.CW-T.WIDE_LEFT-T.GAP, height-T.PAD,
                    body, T.BODY_SM, T.MUTED, anchor=MSO_ANCHOR.MIDDLE)
    elif mode == 'columns':
        width = (T.CW-2*T.GAP)/3
        for i, (head, body) in enumerate(items):
            x = T.MX+i*(width+T.GAP)
            line(s, x, T.STAGE_Y, width,
                 T.GOOD if i == d.get('focus', -1) else T.LINE,
                 T.RULE_THICK if i == d.get('focus', -1) else T.STROKE)
            textbox(s, x, T.STAGE_Y+T.EDITORIAL_INSET,
                    width, T.LEAD_H, f'0{i+1}', T.MID_NUM, T.GOOD, True)
            textbox(s, x, T.STAGE_Y+T.LEAD_H*1.5,
                    width, T.LEAD_H*1.4, head, T.BODY, T.INK, True)
            textbox(s, x, T.STAGE_Y+T.LEAD_H*3,
                    width, T.STAGE_H-T.LEAD_H*3,
                    body, T.BODY_SM, T.MUTED)
    else:
        head, body = items[0]
        rect(s, T.MX, T.STAGE_Y, T.WIDE_LEFT, T.STAGE_H, T.NAVY)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.EDITORIAL_INSET,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET, T.LEAD_H,
                '01', T.MID_NUM, T.WARM, True)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.LEAD_H*1.6,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET, T.LEAD_H*1.5,
                head, T.BODY, T.WHITE, True)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.LEAD_H*3,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET, T.STAGE_H-T.LEAD_H*3,
                body, T.BODY_SM, T.PALE)
        x = T.MX+T.WIDE_LEFT+T.GAP; width = T.CW-T.WIDE_LEFT-T.GAP
        for i, (head, body) in enumerate(items[1:]):
            y = T.STAGE_Y+i*T.STAGE_H/2
            line(s, x, y, width)
            textbox(s, x, y+T.PAD, T.ROW_NUM_W, T.LEAD_H,
                    f'0{i+2}', T.SMALL_NUM, T.GOOD, True)
            textbox(s, x+T.ROW_NUM_W, y+T.PAD,
                    width-T.ROW_NUM_W, T.LEAD_H, head, T.BODY, T.INK, True)
            textbox(s, x+T.ROW_NUM_W, y+T.LEAD_H+T.PAD,
                    width-T.ROW_NUM_W, T.STAGE_H/2-T.LEAD_H-T.PAD,
                    body, T.BODY_SM, T.MUTED)
    bottom(s, d); return s


def compare(prs, d):
    s = base(prs, d); lead(s, d)
    mode = d['_composition']; width = (T.CW-T.GAP)/2
    if mode == 'contrast':
        rect(s, T.MX, T.STAGE_Y, width, T.STAGE_H, T.SURFACE)
        rect(s, T.MX+width+T.GAP, T.STAGE_Y, width, T.STAGE_H, T.NAVY)
    elif mode == 'parallel':
        vertical(s, T.MX+width+T.GAP/2, T.STAGE_Y, T.STAGE_H)
    for i, entry in enumerate((d['left'], d['right'])):
        x = T.MX+i*(width+T.GAP)
        if mode == 'contrast':
            x += T.EDITORIAL_INSET; w = width-2*T.EDITORIAL_INSET
            label_color = T.MUTED if i == 0 else T.WARM
            text_color = T.INK if i == 0 else T.WHITE
            textbox(s, x, T.STAGE_Y+T.EDITORIAL_INSET, w, T.LEAD_H,
                    entry[0], T.BODY, label_color, True)
            line(s, x, T.STAGE_Y+T.LEAD_H+T.EDITORIAL_INSET,
                 w, T.LINE if i == 0 else T.GRAY_DARK)
            textbox(s, x, T.STAGE_Y+T.LEAD_H*1.5+T.EDITORIAL_INSET,
                    w, T.STAGE_H-T.LEAD_H*1.5-T.EDITORIAL_INSET,
                    entry[1], T.BODY_SM, text_color)
        else:
            w = width
            if mode == 'verdict':
                line(s, x, T.STAGE_Y, w,
                     T.MUTED if i == 0 else T.GOOD, T.RULE_THICK)
                textbox(s, x, T.STAGE_Y+T.EDITORIAL_INSET, w, T.LEAD_H,
                        '비교 / 01' if i == 0 else '선택 / 02',
                        T.KICKER, T.MUTED if i == 0 else T.GOOD, True)
                y = T.STAGE_Y+T.LEAD_H*1.2
            else: y = T.STAGE_Y
            textbox(s, x, y, w, T.LEAD_H*1.6,
                    entry[0], T.SMALL_NUM,
                    T.MUTED if i == 0 else T.GOOD, True)
            line(s, x, y+T.LEAD_H*1.7, w)
            textbox(s, x, y+T.LEAD_H*2, w,
                    T.STAGE_BOTTOM-y-T.LEAD_H*2,
                    entry[1], T.BODY_SM, T.INK)
    bottom(s, d); return s


def _widths(d):
    fractions = d.get('widths') or [1/len(d['headers'])]*len(d['headers'])
    total = sum(fractions)
    return [T.CW*f/total for f in fractions]


def table(prs, d):
    s = base(prs, d); lead(s, d)
    mode = d['_composition']; widths = _widths(d)
    rows = [d['headers']]+d['rows']; row_h = T.TABLE_HEIGHT/len(rows)
    if mode == 'matrix':
        for ri, row in enumerate(rows):
            y = T.TABLE_TOP+ri*row_h
            if ri == 0:
                rect(s, T.MX, y, T.CW, row_h, T.SOFT_GREEN)
                line(s, T.MX, y, T.CW, T.GOOD, T.RULE_THICK)
            else:
                rect(s, T.MX, y, widths[0], row_h, T.SURFACE)
                line(s, T.MX, y, T.CW)
            x = T.MX
            for ci, value in enumerate(row):
                textbox(s, x+T.TABLE_PAD, y+T.TABLE_PAD,
                        widths[ci]-2*T.TABLE_PAD, row_h-2*T.TABLE_PAD,
                        value, T.TABLE_COMPACT if len(widths) >= 5 else T.TABLE,
                        T.GOOD if ri == 0 else T.INK,
                        ri == 0 or ci == 0, anchor=MSO_ANCHOR.MIDDLE)
                x += widths[ci]
        line(s, T.MX, T.TABLE_TOP+T.TABLE_HEIGHT, T.CW, T.INK)
    else:
        for ri, row in enumerate(rows):
            y = T.TABLE_TOP+ri*row_h
            if mode == 'bands' and ri > 0 and ri % 2:
                rect(s, T.MX, y, T.CW, row_h, T.SURFACE)
            if ri == 0: line(s, T.MX, y, T.CW, T.INK, T.RULE_THICK)
            elif mode == 'ledger': line(s, T.MX, y, T.CW)
            x = T.MX
            for ci, value in enumerate(row):
                size = T.TABLE_COMPACT if len(str(value)) > 46 else T.TABLE
                textbox(s, x+T.TABLE_PAD, y+T.TABLE_PAD,
                        widths[ci]-2*T.TABLE_PAD, row_h-2*T.TABLE_PAD,
                        value, T.TABLE_HEAD if ri == 0 else size,
                        T.MUTED if ri == 0 else T.INK,
                        ri == 0 or (mode == 'bands' and ci == 0),
                        anchor=MSO_ANCHOR.MIDDLE)
                x += widths[ci]
        line(s, T.MX, T.TABLE_TOP+T.TABLE_HEIGHT, T.CW,
             T.INK if mode == 'bands' else T.LINE)
    bottom(s, d); return s


def chart(prs, d):
    s = base(prs, d); lead(s, d)
    rows = d['rows']
    if d['_composition'] == 'rank':
        sorted_rows = sorted(rows, key=lambda item: item[1], reverse=True)
        first, rest = sorted_rows[0], sorted_rows[1:]
        rect(s, T.MX, T.STAGE_Y, T.WIDE_LEFT, T.STAGE_H, T.NAVY)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.EDITORIAL_INSET,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET, T.LEAD_H,
                'TOP RESPONSE', T.KICKER, T.WARM, True)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.LEAD_H,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET, T.LEAD_H*2,
                f"{first[1]}{d.get('unit', '')}", T.BIG_NUM, T.WHITE, True)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.LEAD_H*3,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET, T.STAGE_H-T.LEAD_H*3,
                first[0], T.BODY, T.PALE, True)
        x = T.MX+T.WIDE_LEFT+T.GAP; w = T.CW-T.WIDE_LEFT-T.GAP
        row_h = T.STAGE_H/max(1, len(rest))
        for i, (label, number) in enumerate(rest):
            y = T.STAGE_Y+i*row_h; line(s, x, y, w)
            textbox(s, x, y+T.PAD, w-T.ROW_NUM_W, row_h-T.PAD,
                    label, T.BODY_SM, T.INK, anchor=MSO_ANCHOR.MIDDLE)
            textbox(s, x+w-T.ROW_NUM_W, y+T.PAD, T.ROW_NUM_W,
                    row_h-T.PAD, str(number), T.BODY, T.GOOD, True,
                    align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    else:
        peak = max((number for _, number in rows), default=1) or 1
        row_h = min(T.CHART_ROW, T.STAGE_H/len(rows))
        gap = min(T.CHART_GAP, (T.STAGE_H-len(rows)*row_h)/max(1, len(rows)-1))
        for i, (label, number) in enumerate(rows):
            y = T.CHART_TOP+i*(row_h+gap)
            textbox(s, T.MX, y, T.CHART_LABEL, row_h,
                    label, T.BODY_SM, anchor=MSO_ANCHOR.MIDDLE)
            bar_x = T.MX+T.CHART_LABEL+T.GAP
            rect(s, bar_x, y+(row_h-T.CHART_BAR_H)/2,
                 T.CHART_BAR, T.CHART_BAR_H, T.PALE)
            if number:
                rect(s, bar_x, y+(row_h-T.CHART_BAR_H)/2,
                     T.CHART_BAR*number/peak, T.CHART_BAR_H,
                     T.GOOD if i == 0 else T.BLUE)
            textbox(s, bar_x+T.CHART_BAR+T.GAP, y,
                    T.CW-T.CHART_LABEL-T.CHART_BAR-2*T.GAP, row_h,
                    f"{number}{d.get('unit', '')}", T.BODY, T.INK, True,
                    anchor=MSO_ANCHOR.MIDDLE)
    bottom(s, d); return s


def process(prs, d):
    s = base(prs, d); lead(s, d)
    items = d['cards']; width = (T.CW-(len(items)-1)*T.GAP)/len(items)
    mode = d['_composition']
    if mode == 'rail':
        rail_y = T.STAGE_Y+T.LEAD_H*1.4
        line(s, T.MX, rail_y, T.CW, T.LINE, T.RULE_THICK)
    for i, (head, body) in enumerate(items):
        x = T.MX+i*(width+T.GAP)
        if mode == 'stair':
            offset = i*T.EDITORIAL_INSET/2
            line(s, x, T.STAGE_Y+offset, width, T.GOOD, T.RULE_THICK)
            textbox(s, x, T.STAGE_Y+offset+T.PAD, width, T.LEAD_H,
                    f'{i+1:02d}', T.MID_NUM, T.GOOD, True)
            y = T.STAGE_Y+offset+T.LEAD_H*1.5
        else:
            rect(s, x, rail_y-T.RULE_THICK*2,
                 T.RULE_THICK*4, T.RULE_THICK*4, T.GOOD)
            textbox(s, x, T.STAGE_Y, width, T.LEAD_H,
                    f'{i+1:02d}', T.SMALL_NUM, T.GOOD, True)
            y = rail_y+T.EDITORIAL_INSET
        textbox(s, x, y, width, T.LEAD_H*1.3, head, T.BODY, T.INK, True)
        textbox(s, x, y+T.LEAD_H*1.5, width,
                T.STAGE_BOTTOM-y-T.LEAD_H*1.5, body, T.BODY_SM, T.MUTED)
    bottom(s, d); return s


def prompt(prs, d):
    s = base(prs, d); lead(s, d)
    items = d['prompt'].split('\n')
    mode = d['_composition']
    left_w = T.CW-T.WIDE_LEFT/2-T.GAP
    if mode == 'annotated':
        try:
            rows = prompt_rows(items, left_w-2*T.EDITORIAL_INSET)
            if sum(height for _, height in rows) > T.STAGE_H-2*T.EDITORIAL_INSET:
                mode = 'source'
        except ValueError:
            mode = 'source'
    if mode == 'source':
        rows = prompt_rows(items, T.PROMPT_BODY_W)
        y = T.STAGE_Y
        for i, (item, row_h) in enumerate(rows):
            line(s, T.MX, y, T.CW)
            textbox(s, T.MX, y+T.PROMPT_PAD_Y,
                    T.PROMPT_LABEL_W-T.TABLE_PAD, row_h-2*T.PROMPT_PAD_Y,
                    f'{i+1:02d}', T.KICKER, T.GOOD, True,
                    anchor=MSO_ANCHOR.TOP)
            textbox(s, T.PROMPT_BODY_X, y+T.PROMPT_PAD_Y,
                    T.PROMPT_BODY_W, row_h-2*T.PROMPT_PAD_Y,
                    item, T.PROMPT_TEXT)
            y += row_h
        line(s, T.MX, y, T.CW)
    else:
        rect(s, T.MX, T.STAGE_Y, left_w, T.STAGE_H, T.SURFACE)
        y = T.STAGE_Y+T.EDITORIAL_INSET
        for i, (item, row_h) in enumerate(rows):
            textbox(s, T.MX+T.EDITORIAL_INSET, y+T.PROMPT_PAD_Y,
                    left_w-2*T.EDITORIAL_INSET, row_h-2*T.PROMPT_PAD_Y,
                    item, T.PROMPT_TEXT)
            y += row_h
            if i < len(rows)-1:
                line(s, T.MX+T.EDITORIAL_INSET, y,
                     left_w-2*T.EDITORIAL_INSET)
        x = T.MX+left_w+T.GAP; w = T.CW-left_w-T.GAP
        textbox(s, x, T.STAGE_Y, w, T.LEAD_H,
                '요청 구조', T.KICKER, T.GOOD, True)
        line(s, x, T.STAGE_Y+T.LEAD_H, w, T.GOOD, T.RULE_THICK)
        picks = items[1:4] if len(items) >= 4 else items[:3]
        for i, item in enumerate(picks):
            stripped = item.strip()
            if d.get('labels'):
                label = d['labels'][i]
            elif stripped.startswith('[') and ']' in stripped:
                label = stripped[1:stripped.index(']')]
            elif ':' in stripped and stripped.index(':') <= 12:
                label = stripped.split(':', 1)[0]
            else:
                label = stripped[:12].rstrip(' .')
            textbox(s, x, T.STAGE_Y+T.LEAD_H*1.4+i*T.LEAD_H*1.2,
                    w, T.LEAD_H, label, T.BODY, T.INK, True)
    bottom(s, d); return s


def exercise(prs, d):
    s = base(prs, d); lead(s, d)
    items = d['steps']; mode = d['_composition']
    if mode == 'worksheet':
        row_h = T.STAGE_H/4
        for i, (head, body) in enumerate(items):
            y = T.STAGE_Y+i*row_h; line(s, T.MX, y, T.CW)
            textbox(s, T.MX, y+T.PAD, T.ROW_NUM_W, row_h-T.PAD,
                    f'0{i+1}', T.SMALL_NUM, T.GOOD, True)
            textbox(s, T.MX+T.ROW_NUM_W, y+T.PAD,
                    T.WIDE_LEFT-T.ROW_NUM_W, row_h-T.PAD,
                    head, T.BODY, T.INK, True, anchor=MSO_ANCHOR.MIDDLE)
            textbox(s, T.MX+T.WIDE_LEFT+T.GAP, y+T.PAD,
                    T.CW-T.WIDE_LEFT-T.GAP, row_h-T.PAD,
                    body, T.BODY_SM, T.MUTED, anchor=MSO_ANCHOR.MIDDLE)
    elif mode == 'path':
        width = (T.CW-3*T.GAP)/4
        line(s, T.MX, T.STAGE_Y+T.LEAD_H*1.3,
             T.CW, T.LINE, T.RULE_THICK)
        for i, (head, body) in enumerate(items):
            x = T.MX+i*(width+T.GAP)
            textbox(s, x, T.STAGE_Y, width, T.LEAD_H,
                    f'{i+1:02d}', T.MID_NUM, T.GOOD, True)
            textbox(s, x, T.STAGE_Y+T.LEAD_H*1.8,
                    width, T.LEAD_H*1.2, head, T.BODY, T.INK, True)
            textbox(s, x, T.STAGE_Y+T.LEAD_H*3.1,
                    width, T.STAGE_H-T.LEAD_H*3.1,
                    body, T.BODY_SM, T.MUTED)
    else:
        width = T.CW-T.WIDE_LEFT-T.GAP
        rect(s, T.MX, T.STAGE_Y, T.WIDE_LEFT, T.STAGE_H, T.SOFT_GREEN)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.EDITORIAL_INSET,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET,
                T.LEAD_H, 'WORKSHEET', T.KICKER, T.GOOD, True)
        textbox(s, T.MX+T.EDITORIAL_INSET, T.STAGE_Y+T.LEAD_H*1.6,
                T.WIDE_LEFT-2*T.EDITORIAL_INSET,
                T.STAGE_H-T.LEAD_H*1.6,
                items[0][0]+'\n\n'+items[0][1], T.BODY, T.INK, True)
        x = T.MX+T.WIDE_LEFT+T.GAP; row_h = T.STAGE_H/3
        for i, (head, body) in enumerate(items[1:]):
            y = T.STAGE_Y+i*row_h; line(s, x, y, width)
            textbox(s, x, y+T.PAD, T.ROW_NUM_W, row_h-T.PAD,
                    f'0{i+2}', T.SMALL_NUM, T.GOOD, True)
            textbox(s, x+T.ROW_NUM_W, y+T.PAD,
                    width-T.ROW_NUM_W, T.LEAD_H, head, T.BODY, T.INK, True)
            textbox(s, x+T.ROW_NUM_W, y+T.LEAD_H+T.PAD,
                    width-T.ROW_NUM_W, row_h-T.LEAD_H-T.PAD,
                    body, T.BODY_SM, T.MUTED)
    bottom(s, d); return s


def picture_fit(s, path, x, y, w, h):
    iw, ih = Image.from_file(path).size
    scale = min(w/iw, h/ih); aw, ah = iw*scale, ih*scale
    s.shapes.add_picture(path, int(x+(w-aw)/2), int(y+(h-ah)/2),
                         width=int(aw), height=int(ah))


def image(prs, d):
    s = base(prs, d)
    path = os.path.join(ASSETS, d['path'])
    rect(s, T.MX, T.STAGE_Y, T.IMAGE_W, T.STAGE_H, T.SURFACE)
    if os.path.exists(path):
        picture_fit(s, path, T.MX+T.EDITORIAL_INSET,
                    T.STAGE_Y+T.EDITORIAL_INSET,
                    T.IMAGE_W-2*T.EDITORIAL_INSET,
                    T.STAGE_H-2*T.EDITORIAL_INSET)
    x = T.MX+T.IMAGE_W+T.GAP; w = T.CW-T.IMAGE_W-T.GAP
    line(s, x, T.STAGE_Y, w, T.GOOD, T.RULE_THICK)
    textbox(s, x, T.STAGE_Y+T.EDITORIAL_INSET, w,
            T.LEAD_H*2, d['lead'], T.LEAD, T.BLUE, True)
    textbox(s, x, T.STAGE_Y+T.LEAD_H*2.3, w,
            T.STAGE_H-T.LEAD_H*2.3, d['body'], T.BODY_SM)
    bottom(s, d); return s


def profile(prs, d):
    s = base(prs, d); lead(s, d)
    left_w = T.CW-T.PHOTO_W-T.GAP
    line(s, T.MX, T.STAGE_Y, left_w, T.GOOD, T.RULE_THICK)
    if d.get('groups'):
        y = T.STAGE_Y+T.EDITORIAL_INSET
        textbox(s, T.MX, y, left_w, T.PROFILE_NAME_H,
                d['name'], T.PROFILE_NAME, T.INK, True)
        y += T.PROFILE_NAME_H+T.PROFILE_GROUP_GAP
        for label, rows in d['groups']:
            textbox(s, T.MX, y, left_w, T.PROFILE_LABEL_H,
                    label, T.KICKER, T.GOOD, True)
            y += T.PROFILE_LABEL_H+T.PROFILE_GROUP_GAP/2
            for row in rows:
                textbox(s, T.MX, y, left_w, T.PROFILE_ROW_H,
                        '• '+row, T.BODY_SM)
                y += T.PROFILE_ROW_H
            y += T.PROFILE_GROUP_GAP
    else:
        textbox(s, T.MX, T.STAGE_Y+T.EDITORIAL_INSET,
                left_w, T.STAGE_H-T.EDITORIAL_INSET, d['body'], T.BODY_SM)
    x = T.MX+left_w+T.GAP
    rect(s, x, T.STAGE_Y, T.PHOTO_W, T.PHOTO_H, T.SURFACE)
    path = os.path.join(ASSETS, d['path'])
    if os.path.exists(path): picture_fit(s, path, x, T.STAGE_Y, T.PHOTO_W, T.PHOTO_H)
    textbox(s, x, T.STAGE_Y+T.PHOTO_H+T.PAD,
            T.PHOTO_W, T.LEAD_H, 'KT 「모두의 AI」 사업 출범식', T.CAPTION, T.MUTED)
    bottom(s, d); return s


def native_table(s, x, y, w, headers, rows, fractions, row_h=T.DOC_ROW_H,
                 head_h=T.DOC_HEAD_H, size=T.DOC_BODY, numeric=(), total=False):
    """숫자 정렬·단위·행의 의미를 보존하는 편집 가능한 데이터 표."""
    tab = s.shapes.add_table(len(rows)+1, len(headers), int(x), int(y),
                             int(w), int(head_h+len(rows)*row_h)).table
    for column, fraction in zip(tab.columns, fractions):
        column.width = int(w*fraction/sum(fractions))
    for ri, values in enumerate([headers]+rows):
        tab.rows[ri].height = int(head_h if ri == 0 else row_h)
        for ci, value in enumerate(values):
            cell = tab.cell(ri, ci)
            cell.margin_left = cell.margin_right = T.TABLE_PAD
            cell.margin_top = cell.margin_bottom = int(T.TABLE_PAD/2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            is_total = total and ri == len(rows)
            cell.fill.fore_color.rgb = T.SOFT_GREEN if is_total else (T.SURFACE if ri == 0 else T.PAPER)
            tcpr = cell._tc.get_or_add_tcPr()
            for name in ('lnL', 'lnR', 'lnT', 'lnB'):
                for old in list(tcpr):
                    if old.tag.endswith('}'+name): tcpr.remove(old)
                border = OxmlElement('a:'+name); border.set('w', str(int(T.STROKE)))
                color = OxmlElement('a:solidFill' if name == 'lnB' else 'a:noFill')
                if name == 'lnB':
                    rgb = OxmlElement('a:srgbClr'); rgb.set('val', str(T.LINE)); color.append(rgb)
                border.append(color); tcpr.append(border)
            tf = cell.text_frame; tf.clear(); tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            for pi, text in enumerate(str(value).split('\n')):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                p.line_spacing = T.BODY_LINES; p.space_before = p.space_after = 0
                p.alignment = PP_ALIGN.RIGHT if ci in numeric else PP_ALIGN.LEFT
                run = p.add_run(); run.text = text
                run.font.name = T.FONT_HEAVY if ri == 0 or is_total else T.FONT
                run.font.size = size; run.font.bold = ri == 0 or is_total
                run.font.color.rgb = T.GOOD if is_total else (T.MUTED if ri == 0 else T.INK)
    return tab


def kpis(s, items, x=T.MX, y=T.STAGE_Y, w=T.CW, h=T.DOC_KPI_H):
    col_w = w/len(items)
    for i, (label, value, context) in enumerate(items):
        xx = x+i*col_w
        if i: vertical(s, xx-T.DOC_INSET, y, h, T.LINE)
        textbox(s, xx, y, col_w-T.DOC_GAP, T.PROFILE_LABEL_H,
                label, T.DOC_LABEL, T.MUTED)
        textbox(s, xx, y+T.PROFILE_LABEL_H+T.PROFILE_GROUP_GAP/2,
                col_w-T.DOC_GAP, T.PROFILE_NAME_H,
                value, T.DOC_VALUE, T.INK, True)
        if context:
            textbox(s, xx, y+h-T.PROFILE_LABEL_H, col_w-T.DOC_GAP,
                    T.PROFILE_LABEL_H, context, T.DOC_LABEL, T.GOOD)


def summary_sheet(prs, d):
    s = base(prs, d); lead(s, d)
    kpis(s, d['metrics'])
    y = T.STAGE_Y+T.DOC_KPI_H+T.DOC_GAP
    native_table(s, T.MX, y, T.CW, d['headers'], d['rows'], d['widths'],
                 row_h=T.DOC_ROW_H, head_h=T.GANTT_HEAD_H,
                 size=T.TABLE_COMPACT, numeric=range(1, len(d['headers'])), total=True)
    textbox(s, T.MX, T.STAGE_BOTTOM-T.PROFILE_ROW_H, T.CW, T.PROFILE_ROW_H,
            d['definition'], T.DOC_SMALL, T.MUTED)
    return s


def report(prs, d):
    s = base(prs, d); lead(s, d)
    rect(s, T.MX, T.STAGE_Y, T.CW, T.DOC_CONCLUSION_H, T.NAVY, rounded=True)
    textbox(s, T.MX+T.DOC_INSET, T.STAGE_Y+T.DOC_INSET,
            T.CW-2*T.DOC_INSET, T.PROFILE_LABEL_H, '결정 요청', T.DOC_LABEL, T.WARM, True)
    textbox(s, T.MX+T.DOC_INSET, T.STAGE_Y+T.PROFILE_LABEL_H+T.DOC_INSET,
            T.CW-2*T.DOC_INSET, T.PROFILE_NAME_H, d['conclusion'], T.BODY, T.WHITE)
    y = T.STAGE_Y+T.DOC_CONCLUSION_H+T.DOC_GAP
    lw = (T.CW-T.DOC_GAP)*T.DOC_LEFT_RATIO
    native_table(s, T.MX, y, lw, d['headers'], d['rows'], d['widths'],
                 numeric=(1, 2), head_h=T.GANTT_HEAD_H)
    x = T.MX+lw+T.DOC_GAP; w = T.CW-lw-T.DOC_GAP
    vertical(s, x, y, T.STAGE_BOTTOM-y-T.DOC_BOTTOM_H)
    x += T.DOC_INSET; w -= T.DOC_INSET
    for i, (label, body) in enumerate(d['actions']):
        yy = y+i*(T.DOC_BLOCK_BODY_H+T.DOC_BLOCK_GAP)
        textbox(s, x, yy, w, T.DOC_BLOCK_TITLE_H, label, T.DOC_BODY, T.GOOD, True)
        textbox(s, x, yy+T.DOC_BLOCK_TITLE_H, w, T.DOC_BLOCK_BODY_H-T.DOC_BLOCK_TITLE_H,
                body, T.DOC_BODY)
    textbox(s, T.MX, T.STAGE_BOTTOM-T.DOC_BOTTOM_H/2, T.CW,
            T.DOC_BOTTOM_H/2, d['definition'], T.DOC_SMALL, T.MUTED)
    return s


def gantt(prs, d):
    s = base(prs, d); lead(s, d)
    meta_w = T.CW*T.GANTT_META_RATIO; week_w = (T.CW-meta_w)/len(d['weeks'])
    widths = [T.CW*f for f in T.GANTT_COLS]
    headers = ['ID', '작업', '담당 역할', '선행', '완료 산출물']
    y = T.STAGE_Y
    rect(s, T.MX, y, T.CW, T.GANTT_HEAD_H, T.SURFACE)
    x = T.MX
    for head, w in zip(headers, widths):
        textbox(s, x+T.TABLE_PAD, y+T.TABLE_PAD, w-2*T.TABLE_PAD,
                T.GANTT_HEAD_H-2*T.TABLE_PAD, head, T.DOC_SMALL, T.MUTED, True,
                anchor=MSO_ANCHOR.MIDDLE)
        x += w
    for i, week in enumerate(d['weeks']):
        textbox(s, T.MX+meta_w+i*week_w, y+T.TABLE_PAD, week_w,
                T.GANTT_HEAD_H-2*T.TABLE_PAD, week, T.DOC_SMALL, T.MUTED, True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for ri, row in enumerate(d['rows']):
        yy = y+T.GANTT_HEAD_H+ri*T.GANTT_ROW_H
        if ri % 2 == 0: rect(s, T.MX, yy, T.CW, T.GANTT_ROW_H, T.SURFACE)
        x = T.MX
        for value, w in zip(row[:5], widths):
            textbox(s, x+T.TABLE_PAD, yy+T.TABLE_PAD, w-2*T.TABLE_PAD,
                    T.GANTT_ROW_H-2*T.TABLE_PAD, value, T.DOC_SMALL,
                    T.INK, anchor=MSO_ANCHOR.MIDDLE)
            x += w
        for wi in range(len(d['weeks'])):
            vertical(s, T.MX+meta_w+wi*week_w, yy, T.GANTT_ROW_H, T.LINE)
        start, end = row[5:7]
        rect(s, T.MX+meta_w+start*week_w+T.TABLE_PAD,
             yy+(T.GANTT_ROW_H-T.GANTT_BAR_H)/2,
             (end-start+1)*week_w-2*T.TABLE_PAD,
             T.GANTT_BAR_H, T.GOOD if ri == len(d['rows'])-1 else T.NAVY, rounded=True)
        line(s, T.MX, yy+T.GANTT_ROW_H, T.CW)
    yy = T.STAGE_BOTTOM-T.DOC_BOTTOM_H
    line(s, T.MX, yy, T.CW, T.GOOD)
    textbox(s, T.MX, yy+T.DOC_INSET, T.CW, T.DOC_BOTTOM_H-T.DOC_INSET,
            d['decision'], T.DOC_BODY)
    return s


def proposal(prs, d):
    s = base(prs, d); lead(s, d)
    rect(s, T.MX, T.STAGE_Y, T.CW, T.DOC_BAND_H, T.SOFT_GREEN, rounded=True)
    textbox(s, T.MX+T.DOC_INSET, T.STAGE_Y+T.DOC_INSET,
            T.CW-2*T.DOC_INSET, T.DOC_BAND_H-2*T.DOC_INSET,
            d['decision'], T.BODY, T.GOOD, True)
    y = T.STAGE_Y+T.DOC_BAND_H+T.DOC_GAP
    w = (T.CW-T.DOC_GAP*2)/2
    vertical(s, T.MX+w+T.DOC_GAP, y, T.STAGE_BOTTOM-y-T.DOC_BOTTOM_H)
    for i, (label, body) in enumerate(d['blocks']):
        x = T.MX+(i%2)*(w+2*T.DOC_GAP)
        yy = y+(i//2)*(T.DOC_BLOCK_BODY_H+T.DOC_BLOCK_GAP)
        textbox(s, x, yy, w, T.DOC_BLOCK_TITLE_H, label, T.DOC_BODY, T.INK, True)
        textbox(s, x, yy+T.DOC_BLOCK_TITLE_H+T.TABLE_PAD,
                w, T.DOC_BLOCK_BODY_H-T.DOC_BLOCK_TITLE_H, body, T.DOC_BODY, T.MUTED)
    line(s, T.MX, T.STAGE_BOTTOM-T.DOC_BOTTOM_H, T.CW)
    textbox(s, T.MX, T.STAGE_BOTTOM-T.DOC_BOTTOM_H+T.DOC_INSET,
            T.CW, T.DOC_BOTTOM_H-T.DOC_INSET, d['close'], T.DOC_BODY)
    return s


def prd(prs, d):
    s = base(prs, d); lead(s, d)
    if d.get('rows'):
        native_table(s, T.MX, T.STAGE_Y, T.CW, d['headers'], d['rows'], d['widths'],
                     row_h=(T.STAGE_H-T.DOC_BOTTOM_H-T.DOC_SPEC_FOOT_GAP-T.GANTT_HEAD_H)/len(d['rows']),
                     head_h=T.GANTT_HEAD_H, size=T.TABLE_COMPACT)
    else:
        width = (T.CW-T.DOC_GAP)/2
        for i, (label, body) in enumerate(d['scope']):
            x = T.MX+i*(width+T.DOC_GAP)
            textbox(s, x, T.STAGE_Y, width, T.DOC_BLOCK_TITLE_H,
                    label, T.DOC_BODY, T.GOOD, True)
            textbox(s, x, T.STAGE_Y+T.DOC_BLOCK_TITLE_H,
                    width, T.DOC_CONCLUSION_H-T.DOC_BLOCK_TITLE_H,
                    body, T.DOC_BODY)
        y = T.STAGE_Y+T.DOC_CONCLUSION_H+T.DOC_GAP
        node_w = (T.CW-(len(d['flow'])-1)*T.DOC_FLOW_ARROW_W)/len(d['flow'])
        for i, label in enumerate(d['flow']):
            x = T.MX+i*(node_w+T.DOC_FLOW_ARROW_W)
            rect(s, x, y, node_w, T.DOC_FLOW_H, T.NAVY, rounded=True)
            textbox(s, x+T.DOC_INSET, y+T.DOC_INSET,
                    node_w-2*T.DOC_INSET, T.DOC_FLOW_H-2*T.DOC_INSET,
                    label, T.DOC_BODY, T.WHITE, True, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE)
            if i < len(d['flow'])-1:
                textbox(s, x+node_w, y+T.DOC_INSET, T.DOC_FLOW_ARROW_W,
                        T.DOC_FLOW_H-2*T.DOC_INSET, '→', T.DOC_BODY, T.GOOD,
                        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        y += T.DOC_FLOW_H+T.DOC_GAP
        for i, (label, body) in enumerate(d['features']):
            x = T.MX+i*(width+T.DOC_GAP)
            line(s, x, y, width)
            textbox(s, x, y+T.DOC_INSET, width, T.DOC_BLOCK_TITLE_H,
                    label, T.DOC_BODY, T.INK, True)
            textbox(s, x, y+T.DOC_INSET+T.DOC_BLOCK_TITLE_H,
                    width, T.DOC_BLOCK_BODY_H-T.DOC_BLOCK_TITLE_H,
                    body, T.DOC_BODY, T.MUTED)
    if d.get('close'):
        textbox(s, T.MX, T.STAGE_BOTTOM-T.DOC_BOTTOM_H+T.DOC_INSET,
                T.CW, T.DOC_BOTTOM_H-T.DOC_INSET, d['close'], T.DOC_BODY, T.GOOD)
    return s


def dashboard(prs, d):
    s = base(prs, d); lead(s, d)
    kpis(s, d['metrics'], h=T.DASH_KPI_H)
    y = T.STAGE_Y+T.DASH_KPI_H+T.DOC_GAP
    lw = (T.CW-T.DOC_GAP)*T.DASH_PLOT_RATIO
    textbox(s, T.MX, y, lw, T.DOC_BLOCK_TITLE_H,
            '제품별 목표 달성률', T.DOC_BODY, T.INK, True)
    by = y+T.DOC_BLOCK_TITLE_H+T.DOC_GAP
    bx = T.MX+T.DASH_LABEL_W
    bw = lw-T.DASH_LABEL_W-T.DASH_VALUE_W-T.DOC_GAP
    target_x = bx+bw/T.DASH_PLOT_MAX
    for i, (label, value) in enumerate(d['bars']):
        yy = by+i*T.DASH_ROW_H
        textbox(s, T.MX, yy, T.DASH_LABEL_W, T.DASH_ROW_H,
                label, T.DOC_BODY, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, bx, yy+(T.DASH_ROW_H-T.DASH_BAR_H)/2, bw,
             T.DASH_BAR_H, T.PALE, rounded=True)
        rect(s, bx, yy+(T.DASH_ROW_H-T.DASH_BAR_H)/2,
             bw*value/T.DASH_PLOT_MAX, T.DASH_BAR_H,
             T.GOOD if value >= 1 else T.NAVY, rounded=True)
        textbox(s, bx+bw+T.TABLE_PAD, yy, T.DASH_VALUE_W,
                T.DASH_ROW_H, f'{value*100:.1f}%', T.DOC_BODY,
                T.INK, True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    vertical(s, target_x, by, len(d['bars'])*T.DASH_ROW_H, T.MUTED, T.STROKE)
    axis_y = by+len(d['bars'])*T.DASH_ROW_H
    textbox(s, bx, axis_y, T.DASH_VALUE_W, T.PROFILE_LABEL_H, '0%', T.DOC_LABEL, T.MUTED)
    textbox(s, target_x-T.DASH_VALUE_W, axis_y, T.DASH_VALUE_W*2,
            T.PROFILE_LABEL_H, '목표 100%', T.DOC_LABEL, T.MUTED, align=PP_ALIGN.CENTER)
    x = T.MX+lw+T.DOC_GAP; w = T.CW-lw-T.DOC_GAP
    vertical(s, x, y, T.DASH_DETAIL_H+T.DOC_GAP)
    x += T.DOC_INSET; w -= T.DOC_INSET
    textbox(s, x, y, w, T.DOC_BLOCK_TITLE_H, '추가 확인', T.DOC_BODY, T.GOOD, True)
    for i, (label, action) in enumerate(d['queue']):
        yy = y+T.DOC_BLOCK_TITLE_H+T.DOC_GAP+i*T.DOC_CONCLUSION_H
        textbox(s, x, yy, w, T.DOC_BLOCK_TITLE_H, label, T.DOC_BODY, T.INK, True)
        textbox(s, x, yy+T.DOC_BLOCK_TITLE_H, w,
                T.DOC_CONCLUSION_H-T.DOC_BLOCK_TITLE_H, action, T.DOC_BODY, T.MUTED)
    textbox(s, T.MX, T.STAGE_BOTTOM-T.PROFILE_ROW_H, T.CW, T.PROFILE_ROW_H,
            d['definition'], T.DOC_SMALL, T.MUTED)
    return s


def pause(prs, d):
    s = base(prs, d, dark=True, title=False)
    line(s, T.MX, T.SECTION_TITLE_Y-T.GAP, T.CW, T.GRAY_DARK)
    textbox(s, T.MX, T.SECTION_TITLE_Y, T.CW, T.LEAD_H*2,
            d['title'], T.COVER, T.WHITE, True)
    textbox(s, T.MX, T.SECTION_SUB_Y+T.LEAD_H, T.CW, T.LEAD_H*2,
            d.get('lead', ''), T.BODY, T.PALE)
    return s


TYPES = {name: globals()[name] for name in
         'cover section triad compare chart table process prompt exercise image profile summary_sheet report gantt proposal prd dashboard pause'.split()}
