"""140장 PDF의 콘택트시트와 PPTX 기본 품질 보고."""
import os
import sys
from collections import Counter
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from slides_data import SLIDES

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
PPTX = os.path.join(ROOT, 'output', 'daeryzen-ai-office-20261007.pptx')
PAGES = os.path.join(ROOT, 'output', 'qa-pages')
OUT = os.path.join(ROOT, 'output', 'qa-contact')

def inspect():
    prs = Presentation(PPTX)
    assert len(prs.slides) == len(SLIDES)
    titles = []
    issues = []
    sizes = []
    for index, slide in enumerate(prs.slides, 1):
        title = SLIDES[index-1]['title']
        slide_text = '\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame)
        titles.append(title)
        if not title or title not in slide_text:
            issues.append(f'{index}: 제목 누락')
        for sh in slide.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs.slide_width + 1000 or sh.top + sh.height > prs.slide_height + 1000:
                issues.append(f'{index}: 슬라이드 밖 도형')
            if sh.has_text_frame:
                for paragraph in sh.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            sizes.append(run.font.size.pt)
    duplicates = [k for k, n in Counter(titles).items() if n > 1]
    if duplicates:
        issues.append(f'제목 중복: {duplicates}')
    print(f'slides={len(prs.slides)}; min_font={min(sizes):.1f}pt; duplicate_titles={len(duplicates)}; issues={len(issues)}')
    for issue in issues:
        print(issue)

def contact():
    os.makedirs(OUT, exist_ok=True)
    paths = sorted(os.path.join(PAGES, name) for name in os.listdir(PAGES)
                   if name.startswith('page-') and name.endswith('.jpg'))
    assert len(paths) == len(SLIDES), len(paths)
    cell_w = 400
    cell_h = 270
    for group in range((len(paths)+19)//20):
        canvas = Image.new('RGB', (cell_w*5, cell_h*4), 'white')
        draw = ImageDraw.Draw(canvas)
        for slot, path in enumerate(paths[group*20:(group+1)*20]):
            image = Image.open(path).convert('RGB')
            image.thumbnail((cell_w-12, cell_h-28))
            x = (slot%5)*cell_w+6
            y = (slot//5)*cell_h+20
            canvas.paste(image, (x, y))
            draw.text((x, y-17), f'{group*20+slot+1:03d}', fill='black')
        canvas.save(os.path.join(OUT, f'contact-{group+1:02d}.jpg'), quality=84)
    print(f'contact_sheets={group+1}')

if __name__ == '__main__':
    inspect()
    contact()
